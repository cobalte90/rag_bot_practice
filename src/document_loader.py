"""Document loading utilities."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx"}


class DocumentLoadingError(RuntimeError):
    """Raised when source documents cannot be loaded."""


@dataclass(slots=True)
class DocumentPage:
    """A single extracted page or logical page."""

    page_number: int
    text: str


@dataclass(slots=True)
class LoadedDocument:
    """A loaded source document."""

    source_path: Path
    source_file: str
    pages: list[DocumentPage]


def discover_documents(raw_data_dir: Path) -> list[Path]:
    """Return supported source documents from the raw data directory."""

    if not raw_data_dir.exists():
        raise DocumentLoadingError(
            f"Raw data directory does not exist: {raw_data_dir}. "
            "Create it and place `postgresql-18-US.pdf` inside."
        )

    documents = sorted(
        path for path in raw_data_dir.iterdir() if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )
    if not documents:
        raise DocumentLoadingError(
            f"No supported documents found in {raw_data_dir}. "
            "Supported formats: .pdf, .txt, .docx"
        )
    return documents


def load_documents(raw_data_dir: Path) -> list[LoadedDocument]:
    """Load all supported documents from the raw data directory."""

    return [load_document(path) for path in discover_documents(raw_data_dir)]


def load_document(path: Path) -> LoadedDocument:
    """Load a single document based on its file extension."""

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".txt":
        return _load_txt(path)
    if suffix == ".docx":
        return _load_docx(path)
    raise DocumentLoadingError(f"Unsupported file type: {path.suffix}")


def _load_pdf(path: Path) -> LoadedDocument:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # pragma: no cover - library-specific errors are varied
        raise DocumentLoadingError(f"Failed to read PDF file {path}: {exc}") from exc

    pages: list[DocumentPage] = []
    for page_index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # pragma: no cover - library-specific errors are varied
            raise DocumentLoadingError(
                f"Failed to extract text from PDF page {page_index} in {path}: {exc}"
            ) from exc
        pages.append(DocumentPage(page_number=page_index, text=text))

    return LoadedDocument(source_path=path, source_file=path.name, pages=pages)


def _load_txt(path: Path) -> LoadedDocument:
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = path.read_text(encoding="utf-8-sig")
    except Exception as exc:  # pragma: no cover - filesystem-specific
        raise DocumentLoadingError(f"Failed to read text file {path}: {exc}") from exc

    return LoadedDocument(
        source_path=path,
        source_file=path.name,
        pages=[DocumentPage(page_number=1, text=text)],
    )


def _load_docx(path: Path) -> LoadedDocument:
    try:
        document = DocxDocument(str(path))
    except Exception as exc:  # pragma: no cover - library-specific
        raise DocumentLoadingError(f"Failed to read DOCX file {path}: {exc}") from exc

    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    return LoadedDocument(
        source_path=path,
        source_file=path.name,
        pages=[DocumentPage(page_number=1, text=text)],
    )
